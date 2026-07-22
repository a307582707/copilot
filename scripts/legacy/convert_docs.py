#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Convert project markdown artifacts to DOCX and Mermaid diagrams to PNG.

Usage examples:
  python3 scripts/legacy/convert_docs.py md-to-docx "in.md" "out.docx"
  python3 scripts/legacy/convert_docs.py mermaid-md-to-png "diagram.md" "diagram.png"
"""

import argparse
import base64
import json
import os
import re
import sys
import urllib.request
import zlib
from typing import List, Optional

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
except ImportError:
    # Fallback or exit if python-docx not installed
    pass


def _set_default_style(doc):
    # Best-effort: make Chinese readable when opened on Windows.
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    # East Asia font
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)


def _add_inline_md_runs(p, text: str) -> None:
    """
    Minimal inline Markdown:
    - **bold**
    - `code`
    Keeps other markers as-is if unbalanced.
    """
    def add_normal(seg: str) -> None:
        if not seg:
            return
        # parse **bold** inside normal segments
        pos = 0
        for m in re.finditer(r"\*\*(.+?)\*\*", seg):
            if m.start() > pos:
                p.add_run(seg[pos:m.start()])
            r = p.add_run(m.group(1))
            r.bold = True
            pos = m.end()
        if pos < len(seg):
            p.add_run(seg[pos:])

    # Split by backticks, alternating normal/code segments
    parts = text.split("`")
    if len(parts) == 1:
        add_normal(text)
        return
    if len(parts) % 2 == 0:
        # unbalanced backticks, keep literal
        add_normal(text)
        return

    for idx, part in enumerate(parts):
        if idx % 2 == 0:
            add_normal(part)
        else:
            r = p.add_run(part)
            r.font.name = "Courier New"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _is_table_sep(line: str) -> bool:
    # Markdown table separator row: | --- | :---: | ---: |
    s = line.strip()
    if "|" not in s:
        return False
    s = s.strip("|").strip()
    if not s:
        return False
    cells = [c.strip() for c in s.split("|")]
    if not cells:
        return False
    for c in cells:
        if not c:
            return False
        if not re.fullmatch(r":?-{3,}:?", c):
            return False
    return True


def _parse_table_row(line: str) -> List[str]:
    s = line.rstrip("\n").strip()
    if "|" not in s:
        return []
    cells = [c.strip() for c in s.strip().strip("|").split("|")]
    return cells


def md_to_docx(in_md: str, out_docx: str) -> None:
    with open(in_md, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    _set_default_style(doc)

    i = 0
    in_code = False
    code_lang = None
    code_buf = []  # type: List[str]

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Fenced code blocks
        if stripped.startswith("```"):
            fence = stripped
            if not in_code:
                in_code = True
                code_lang = fence[3:].strip() or None
                code_buf = []
            else:
                # close
                in_code = False
                # Render code block
                if code_lang:
                    p = doc.add_paragraph("```" + code_lang)
                    p.style = doc.styles["No Spacing"] if "No Spacing" in doc.styles else doc.styles["Normal"]
                for cl in code_buf:
                    p = doc.add_paragraph(cl)
                    p.style = doc.styles["No Spacing"] if "No Spacing" in doc.styles else doc.styles["Normal"]
                    for r in p.runs:
                        r.font.name = "Courier New"
                        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                p = doc.add_paragraph("```")
                p.style = doc.styles["No Spacing"] if "No Spacing" in doc.styles else doc.styles["Normal"]
                code_lang = None
                code_buf = []
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # Horizontal rule
        if stripped == "---":
            doc.add_paragraph("—" * 16)
            i += 1
            continue

        # Headings
        if stripped.startswith("#"):
            m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if m:
                level = len(m.group(1))
                text = m.group(2).strip()
                # Map: '#' -> Title-ish, '##' -> H1, '###' -> H2...
                doc_level = max(0, min(9, level - 1))
                doc.add_heading(text, level=doc_level)
                i += 1
                continue

        # Blockquote
        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            p = doc.add_paragraph("")
            _add_inline_md_runs(p, text)
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # Markdown table
        if "|" in stripped and (i + 1) < len(lines) and _is_table_sep(lines[i + 1]):
            header = _parse_table_row(lines[i])
            i += 2  # skip sep row
            rows = []  # type: List[List[str]]
            while i < len(lines) and "|" in lines[i].strip():
                row = _parse_table_row(lines[i])
                # Stop on empty-ish row
                if not any(c.strip() for c in row):
                    break
                rows.append(row)
                i += 1

            col_count = max(1, len(header))
            t = doc.add_table(rows=1, cols=col_count)
            if "Table Grid" in doc.styles:
                t.style = "Table Grid"

            # Header row
            hdr_cells = t.rows[0].cells
            for ci in range(col_count):
                text = header[ci] if ci < len(header) else ""
                hdr_cells[ci].text = ""
                hp = hdr_cells[ci].paragraphs[0]
                _add_inline_md_runs(hp, text)
                for r in hp.runs:
                    r.bold = True

            # Data rows
            for row in rows:
                tr = t.add_row().cells
                for ci in range(col_count):
                    cell_text = row[ci] if ci < len(row) else ""
                    tr[ci].text = ""
                    cp = tr[ci].paragraphs[0]
                    _add_inline_md_runs(cp, cell_text)
            continue

        # Bullet list
        if stripped.startswith("- ") or stripped.startswith("* "):
            text = stripped[2:].strip()
            p = doc.add_paragraph("")
            _add_inline_md_runs(p, text)
            if "List Bullet" in doc.styles:
                p.style = "List Bullet"
            i += 1
            continue

        # Blank line
        if stripped == "":
            doc.add_paragraph("")
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph("")
        _add_inline_md_runs(p, line)
        i += 1

    doc.save(out_docx)


def _extract_mermaid_code(md_path: str) -> str:
    with open(md_path, "r", encoding="utf-8") as f:
        text = f.read()

    m = re.search(r"```mermaid\s*(.*?)\s*```", text, flags=re.DOTALL)
    if not m:
        raise RuntimeError("No mermaid code block found in: %s" % md_path)
    return m.group(1).strip() + "\n"


def _b64url_no_pad(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).decode("ascii").rstrip("=")


def _encode_mermaid_pako(code: str, theme: str) -> str:
    """
    Encode Mermaid payload as mermaid.ink expects:
    a zlib-compressed JSON object (mermaid.live style), base64url, prefixed with 'pako:'.
    This is more robust than plain base64 for large diagrams because the URL path becomes much shorter.
    """
    payload = {
        "code": code,
        # mermaid.ink expects the 'mermaid' field to be a JSON string
        "mermaid": json.dumps({"theme": theme}),
        "autoSync": True,
        "updateDiagram": True,
    }
    raw = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    z = zlib.compress(raw, level=9)
    return "pako:" + _b64url_no_pad(z)


def _fetch_png_from_mermaid_ink(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "codesprite-docs/1.0"})
    with urllib.request.urlopen(req, timeout=30) as r:
        return r.read()


def mermaid_md_to_png(
    in_md: str,
    out_png: str,
    theme: str = "neutral",
    bg_color: str = "white",
    scale: int = 1,
    width: Optional[int] = None,
    height: Optional[int] = None,
) -> None:
    code = _extract_mermaid_code(in_md)
    b64 = _b64url_no_pad(code.encode("utf-8"))
    url_plain = "https://mermaid.ink/img/{b}?type=png&theme={theme}&bgColor={bg}".format(
        b=b64, theme=theme, bg=bg_color
    )
    url_pako = "https://mermaid.ink/img/{b}?type=png&theme={theme}&bgColor={bg}".format(
        b=_encode_mermaid_pako(code, theme=theme), theme=theme, bg=bg_color
    )

    # mermaid.ink: scale requires width or height
    q_extra = ""
    if width is not None:
        q_extra += "&width=%d" % int(width)
    if height is not None:
        q_extra += "&height=%d" % int(height)
    if scale and int(scale) != 1:
        if width is None and height is None:
            q_extra += "&width=1904"
        q_extra += "&scale=%d" % int(scale)
    if q_extra:
        url_plain += q_extra
        url_pako += q_extra

    # Some proxies/servers reject overly long URLs (HTTP 400). Prefer plain for small
    # diagrams, fallback to pako for large ones or on HTTP errors.
    data: bytes
    try:
        if len(url_plain) <= 7000:
            data = _fetch_png_from_mermaid_ink(url_plain)
        else:
            data = _fetch_png_from_mermaid_ink(url_pako)
    except Exception:
        data = _fetch_png_from_mermaid_ink(url_pako)

    if not data or data[:8] != b"\x89PNG\r\n\x1a\n":
        raise RuntimeError("Unexpected response (not PNG) from mermaid.ink")

    with open(out_png, "wb") as f:
        f.write(data)


def main(argv: List[str]) -> int:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd")

    ap_md = sub.add_parser("md-to-docx")
    ap_md.add_argument("in_md")
    ap_md.add_argument("out_docx")

    ap_mm = sub.add_parser("mermaid-md-to-png")
    ap_mm.add_argument("in_md")
    ap_mm.add_argument("out_png")
    ap_mm.add_argument("--theme", default="neutral")
    ap_mm.add_argument("--bg-color", default="white")
    ap_mm.add_argument("--scale", type=int, default=1)
    ap_mm.add_argument("--width", type=int, default=None)
    ap_mm.add_argument("--height", type=int, default=None)

    args = ap.parse_args(argv)
    if not getattr(args, "cmd", None):
        ap.print_help()
        return 2

    if args.cmd == "md-to-docx":
        md_to_docx(args.in_md, args.out_docx)
        return 0
    if args.cmd == "mermaid-md-to-png":
        mermaid_md_to_png(
            args.in_md,
            args.out_png,
            theme=args.theme,
            bg_color=args.bg_color,
            scale=args.scale,
            width=args.width,
            height=args.height,
        )
        return 0

    ap.error("Unknown cmd: %r" % (args.cmd,))
    return 2


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
